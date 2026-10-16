import datetime
from decimal import Decimal
from io import BytesIO
from unittest import mock

from django.core.cache import cache
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.throttling import ScopedRateThrottle

from apps.activity_log.models import ActivityLog
from apps.common.enums import Status
from apps.examination.models import Examination
from apps.invoice.enums import InvoiceStage
from apps.invoice.models import Invoice
from apps.order.enums import OrderStage
from apps.order.models import Order
from apps.order.services import OrderService
from apps.order.tasks import check_order_stage
from apps.third_party.qpay.services import QPayService
from apps.third_party.qpay.tests import invoice_response
from helpers.tests_factories import AdminAPITestCase, make_examination, make_invoice, make_order, make_user

ORDERS_URL = "/orders/"


def order_body(examination, **fields):
    data = {"firstName": "Tuya", "lastName": "Gan", "phone": "99112233", "register": "УБ90010112",
            "examination": examination.id}
    data.update(fields)
    return data


def set_created(order, days_ago):
    Order.objects.filter(pk=order.pk).update(created_at=timezone.now() - datetime.timedelta(days=days_ago))


class OrderTestCase(AdminAPITestCase):
    def setUp(self):
        super().setUp()
        self.examination = make_examination()

    def assert_locked(self, examination, expected):
        examination.refresh_from_db()
        self.assertIs(examination.is_lock, expected)

    def assert_stage(self, order, expected):
        order.refresh_from_db()
        self.assertEqual(order.stage, expected)


@mock.patch.object(check_order_stage, "apply_async")
@mock.patch.object(QPayService, "create_invoice")
class OrderCreateTests(OrderTestCase):
    def setUp(self):
        super().setUp()
        self.client.force_authenticate(None)

    def test_public_create_creates_invoice_locks_slot_and_schedules_expiry(self, create_invoice, apply_async):
        create_invoice.return_value = (invoice_response("inv-5"), None)
        before = timezone.now()
        with self.captureOnCommitCallbacks(execute=True):
            response = self.client.post(ORDERS_URL, order_body(self.examination, isManual=True), format="json")
            apply_async.assert_not_called()
        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.content)
        body = response.json()
        order = Order.objects.get(pk=body["id"])
        invoice = Invoice.objects.get(order=order)
        self.assertEqual((order.stage, order.is_manual), (OrderStage.PENDING, False))
        self.assertEqual(body["invoice"]["qpay"]["invoiceId"], "inv-5")
        self.assertEqual(body["examinationDepositAmount"], 10000.0)
        self.assertEqual(invoice.amount, Decimal("10000.00"))
        self.assert_locked(self.examination, True)
        apply_async.assert_called_once_with(args=(order.id, invoice.id), eta=order.expire_at)
        self.assertAlmostEqual((order.expire_at - before).total_seconds(), 180, delta=10)
        self.assertTrue(ActivityLog.objects.filter(object_id=order.id, action="CREATE").exists())

    def test_invalid_requests_are_rejected(self, create_invoice, apply_async):
        locked = make_examination(is_lock=True)
        cases = [
            (order_body(self.examination, phone="12ab"), "phone"),
            ({"firstName": "Tuya", "examination": self.examination.id}, "register"),
            (order_body(locked), "examination"),
        ]
        for data, field in cases:
            with self.subTest(field=field):
                response = self.client.post(ORDERS_URL, data, format="json")
                self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
                self.assertIn(field, response.json())
        create_invoice.assert_not_called()
        apply_async.assert_not_called()

    def test_qpay_failure_rolls_back_order_invoice_and_lock(self, create_invoice, apply_async):
        create_invoice.return_value = (None, "qpay down")
        with self.captureOnCommitCallbacks(execute=True) as callbacks:
            response = self.client.post(ORDERS_URL, order_body(self.examination), format="json")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("invoice", response.json())
        self.assertFalse(Order.objects.exists())
        self.assertFalse(Invoice.objects.exists())
        self.assertFalse(ActivityLog.objects.exists())
        self.assert_locked(self.examination, False)
        self.assertEqual(callbacks, [])
        apply_async.assert_not_called()


class PublicOrderLookupTests(OrderTestCase):
    """The patient app's "my orders" page lists paid orders anonymously by exact phone or register."""

    PUBLIC_FIELDS = {"id", "createdAt", "examinationTypeName", "examinationDate", "examinationTime",
                     "examinationTotalAmount", "examinationDepositAmount"}

    def setUp(self):
        super().setUp()
        self.paid = make_order(stage=OrderStage.PAID, phone="99887766", register="УБ90010199")
        self.pending = make_order(stage=OrderStage.PENDING, phone="99887766")
        make_order(stage=OrderStage.PAID, phone="88776655", register="УБ80010188")
        self.client.force_authenticate(None)

    def lookup(self, search):
        return self.client.get(ORDERS_URL, {"search": search, "stage": OrderStage.PAID})

    def test_exact_phone_or_register_returns_only_paid_orders_without_personal_data(self):
        for search in ("99887766", "УБ90010199", "уб90010199"):
            with self.subTest(search=search):
                response = self.lookup(search)
                self.assertEqual(response.status_code, status.HTTP_200_OK)
                rows = self.results(response)
                self.assertEqual([row["id"] for row in rows], [self.paid.id])
                self.assertEqual(set(rows[0]), self.PUBLIC_FIELDS)

    def test_partial_or_missing_search_returns_nothing(self):
        for params in ({"search": "9988"}, {"search": "УБ9001"}, {}):
            with self.subTest(params=params):
                response = self.client.get(ORDERS_URL, params)
                self.assertEqual(response.status_code, status.HTTP_200_OK)
                self.assertEqual(self.results(response), [])

    def test_staff_list_keeps_full_rows_and_partial_search(self):
        self.client.force_authenticate(self.admin)
        rows = self.results(self.client.get(ORDERS_URL, {"search": "9988"}))
        self.assertEqual({row["id"] for row in rows}, {self.paid.id, self.pending.id})
        self.assertIn("register", rows[0])

    def test_lookup_is_rate_limited(self):
        cache.clear()
        with mock.patch.dict(ScopedRateThrottle.THROTTLE_RATES, {"order_lookup": "2/minute"}):
            codes = [self.lookup("99887766").status_code for _ in range(3)]
        self.assertEqual(codes, [200, 200, 429])


class OrderPermissionTests(OrderTestCase):
    def test_anonymous_can_only_create(self):
        order = make_order(examination=self.examination)
        self.client.force_authenticate(None)
        detail = f"{ORDERS_URL}{order.id}/"
        requests = [
            ("get", detail), ("get", f"{ORDERS_URL}export/"),
            ("get", f"{ORDERS_URL}report/"), ("post", f"{ORDERS_URL}manual/"), ("put", detail),
            ("patch", detail), ("patch", f"{detail}approve/"), ("patch", f"{detail}cancel/"), ("delete", detail),
        ]
        for method, url in requests:
            with self.subTest(method=method, url=url):
                response = getattr(self.client, method)(url, {}, format="json")
                self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)
                self.assertNotIn(order.register, response.content.decode())

    def test_any_logged_in_staff_account_can_manage_orders(self):
        order = make_order(examination=self.examination, stage=OrderStage.CANCELLED)
        self.client.force_authenticate(make_user())
        response = self.client.patch(f"{ORDERS_URL}{order.id}/approve/", {}, format="json")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assert_stage(order, OrderStage.PAID)


class OrderManualAndListTests(OrderTestCase):
    def test_manual_create_is_paid_and_locks_examination(self):
        response = self.client.post(f"{ORDERS_URL}manual/", order_body(self.examination), format="json")
        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.content)
        order = Order.objects.get(pk=response.json()["id"])
        self.assertEqual((order.stage, order.is_manual, order.created_by), (OrderStage.PAID, True, self.admin))
        self.assert_locked(self.examination, True)

    def test_manual_create_on_locked_examination_saves_nothing(self):
        locked = make_examination(is_lock=True)
        response = self.client.post(f"{ORDERS_URL}manual/", order_body(locked), format="json")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertFalse(Order.objects.exists())

    def test_manual_create_requires_authentication(self):
        self.client.force_authenticate(None)
        response = self.client.post(f"{ORDERS_URL}manual/", order_body(self.examination), format="json")
        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)

    def test_list_returns_active_orders_with_totals_and_search(self):
        make_order(examination=self.examination, phone="99000001")
        make_order(phone="88000002")
        make_order(status=Status.DELETED)
        response = self.client.get(ORDERS_URL)
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        body = response.json()
        self.assertEqual(body["count"], 2)
        self.assertEqual((body["totalAmountSum"], body["depositAmountSum"]), (100000.0, 20000.0))
        found = self.results(self.client.get(ORDERS_URL, {"search": "99000"}))
        self.assertEqual([o["phone"] for o in found], ["99000001"])
        self.assertEqual(found[0]["examinationTypeName"], "Ultrasound")

    def test_delete_is_disabled(self):
        order = make_order()
        response = self.client.delete(f"{ORDERS_URL}{order.id}/")
        self.assertEqual(response.status_code, status.HTTP_405_METHOD_NOT_ALLOWED)
        OrderService.delete(order.id)
        order.refresh_from_db()
        self.assertEqual(order.status, Status.DELETED)


class OrderUpdateTests(OrderTestCase):
    def setUp(self):
        super().setUp()
        self.order = make_order(examination=self.examination)
        Examination.objects.filter(pk=self.examination.pk).update(is_lock=True)
        self.url = f"{ORDERS_URL}{self.order.id}/"

    def put(self, examination):
        body = order_body(examination, firstName="Saraa", reason="moved", isRefund=False)
        return self.client.put(self.url, body, format="json")

    def test_moving_to_free_examination_swaps_locks(self):
        target = make_examination()
        response = self.put(target)
        self.assertEqual(response.status_code, status.HTTP_202_ACCEPTED)
        self.assertEqual(response.json()["firstName"], "Saraa")
        self.order.refresh_from_db()
        self.assertEqual((self.order.examination, self.order.updated_by), (target, self.admin))
        self.assert_locked(target, True)
        self.assert_locked(self.examination, False)

    def test_moving_to_locked_examination_is_rejected(self):
        response = self.put(make_examination(is_lock=True))
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.order.refresh_from_db()
        self.assertEqual(self.order.examination, self.examination)

    def test_missing_or_invalid_fields_return_400(self):
        missing_examination = order_body(self.examination)
        missing_examination.pop("examination")
        missing_name = order_body(self.examination)
        missing_name.pop("firstName")
        unknown_examination = order_body(self.examination)
        unknown_examination["examination"] = 999999
        cases = [
            (missing_examination, "examination"),
            (missing_name, "firstName"),
            (order_body(self.examination, phone="12ab"), "phone"),
            (unknown_examination, "examination"),
        ]
        for body, field in cases:
            with self.subTest(field=field):
                response = self.client.put(self.url, body, format="json")
                self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
                self.assertIn(field, response.json())
        self.order.refresh_from_db()
        self.assertEqual((self.order.first_name, self.order.examination), ("Tuya", self.examination))

    def test_update_keeps_optional_notes_and_handles_order_without_examination(self):
        Order.objects.filter(pk=self.order.pk).update(examination=None, reason="keep", is_refund=True)
        target = make_examination()
        response = self.client.put(self.url, order_body(target), format="json")
        self.assertEqual(response.status_code, status.HTTP_202_ACCEPTED, response.content)
        self.order.refresh_from_db()
        self.assertEqual((self.order.examination, self.order.reason, self.order.is_refund), (target, "keep", True))
        self.assert_locked(target, True)


class OrderStageActionTests(OrderTestCase):
    reason = {"reason": "client request", "isRefund": True}

    def test_cancel_paid_order_unlocks_examination(self):
        order = make_order(examination=make_examination(is_lock=True), stage=OrderStage.PAID)
        response = self.client.patch(f"{ORDERS_URL}{order.id}/cancel/", self.reason, format="json")
        self.assertEqual(response.json(), {"message": "Order cancelled"})
        self.assert_stage(order, OrderStage.CANCELLED)
        self.assertEqual((order.reason, order.is_refund, order.updated_by), ("client request", True, self.admin))
        self.assert_locked(order.examination, False)

    def test_approve_cancelled_order_locks_examination(self):
        order = make_order(examination=self.examination, stage=OrderStage.CANCELLED)
        response = self.client.patch(f"{ORDERS_URL}{order.id}/approve/", self.reason, format="json")
        self.assertEqual(response.json(), {"message": "Order approved"})
        self.assert_stage(order, OrderStage.PAID)
        self.assert_locked(self.examination, True)

    def test_empty_body_changes_only_the_stage(self):
        for action, start, end in (("cancel", OrderStage.PAID, OrderStage.CANCELLED),
                                   ("approve", OrderStage.CANCELLED, OrderStage.PAID)):
            with self.subTest(action=action):
                order = make_order(examination=make_examination(is_lock=start == OrderStage.PAID), stage=start,
                                   reason="old", is_refund=True)
                response = self.client.patch(f"{ORDERS_URL}{order.id}/{action}/", {}, format="json")
                self.assertEqual(response.status_code, status.HTTP_200_OK, response.content)
                self.assert_stage(order, end)
                self.assertEqual((order.reason, order.is_refund, order.updated_by), ("old", True, self.admin))

    def test_invalid_body_and_missing_examination_return_400(self):
        cancelled = make_order(examination=None, stage=OrderStage.CANCELLED)
        paid = make_order(stage=OrderStage.PAID)
        cases = [(cancelled, "approve", {}, "examination"), (paid, "cancel", {"isRefund": "maybe"}, "isRefund")]
        for order, action, body, field in cases:
            with self.subTest(action=action):
                response = self.client.patch(f"{ORDERS_URL}{order.id}/{action}/", body, format="json")
                self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
                self.assertIn(field, response.json())
        self.assert_stage(cancelled, OrderStage.CANCELLED)
        self.assert_stage(paid, OrderStage.PAID)

    def test_cancel_order_without_examination(self):
        order = make_order(examination=None, stage=OrderStage.PAID)
        response = self.client.patch(f"{ORDERS_URL}{order.id}/cancel/", {}, format="json")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assert_stage(order, OrderStage.CANCELLED)

    def test_invalid_transitions_are_rejected(self):
        cases = [
            ("approve", OrderStage.PAID, False),
            ("approve", OrderStage.CANCELLED, True),
            ("cancel", OrderStage.PENDING, False),
            ("cancel", OrderStage.CANCELLED, False),
        ]
        for action, stage, locked in cases:
            with self.subTest(action=action, stage=stage, locked=locked):
                order = make_order(examination=make_examination(is_lock=locked), stage=stage)
                response = self.client.patch(f"{ORDERS_URL}{order.id}/{action}/", self.reason, format="json")
                self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
                self.assert_stage(order, stage)


class OrderReportExportTests(OrderTestCase):
    def setUp(self):
        super().setUp()
        self.today = timezone.localdate().isoformat()
        self.paid = make_order(stage=OrderStage.PAID, is_manual=True, register="")
        self.cancelled = make_order(stage=OrderStage.CANCELLED, is_refund=True)
        self.pending = make_order()
        set_created(make_order(stage=OrderStage.PAID), days_ago=30)

    def test_date_parameters_are_validated(self):
        for path in ("report", "export"):
            for params in ({}, {"start_date": self.today}, {"start_date": "2030/01/01", "end_date": self.today}):
                with self.subTest(path=path, params=params):
                    response = self.client.get(f"{ORDERS_URL}{path}/", params)
                    self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_report_groups_by_stage(self):
        response = self.client.get(f"{ORDERS_URL}report/", {"start_date": self.today, "end_date": self.today})
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        rows = {row["stage"]: row for row in response.json()["results"]}
        self.assertEqual({stage: row["count"] for stage, row in rows.items()},
                         {OrderStage.PAID: 1, OrderStage.CANCELLED: 1, OrderStage.PENDING: 1})
        self.assertEqual(rows[OrderStage.PAID]["depositAmount"], 10000.0)
        self.assertIsNone(rows[OrderStage.PENDING]["depositAmount"])

    def test_export_returns_xlsx_with_paid_and_cancelled_orders(self):
        response = self.client.get(f"{ORDERS_URL}export/", {"start_date": self.today, "end_date": self.today})
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertIn(f'filename="orders_{self.today}_{self.today}.xlsx"', response["Content-Disposition"])
        sheet = load_workbook(BytesIO(response.content)).active
        rows = list(sheet.iter_rows(values_only=True))
        self.assertEqual(sheet.title, "Захиалга")
        self.assertEqual((rows[0][0], rows[0][-1], len(rows[0])), ("ID", "Шинэчилсэн огноо", 15))
        self.assertEqual([r[0] for r in rows[1:]], [self.cancelled.id, self.paid.id])
        cancelled_row, paid_row = rows[1], rows[2]
        self.assertEqual(cancelled_row[10:13], ("Цуцалсан", "Үгүй", "Тийм"))
        self.assertEqual(paid_row[10:13], ("Төлөгдсөн", "Тийм", "Үгүй"))
        self.assertIsNone(paid_row[3])
        self.assertEqual((paid_row[6], paid_row[7], paid_row[8]), ("Ultrasound", "2030-01-10 09:00:00", 10000))


class CheckOrderStageTaskTests(OrderTestCase):
    def setUp(self):
        super().setUp()
        Examination.objects.filter(pk=self.examination.pk).update(is_lock=True)
        self.order = make_order(examination=self.examination)
        self.invoice = make_invoice(order=self.order)

    def expire(self, minutes):
        Order.objects.filter(pk=self.order.pk).update(expire_at=timezone.now() + datetime.timedelta(minutes=minutes))

    def test_expired_pending_order_is_cancelled(self):
        self.expire(-1)
        check_order_stage(self.order.id, self.invoice.id)
        self.assert_stage(self.order, OrderStage.CANCELLED)
        self.invoice.refresh_from_db()
        self.assertEqual(self.invoice.stage, InvoiceStage.CANCELLED)
        self.assert_locked(self.examination, False)

    def set_stage(self, stage):
        Order.objects.filter(pk=self.order.pk).update(stage=stage)

    def test_orders_that_are_final_or_not_expired_are_left_alone(self):
        for stage, minutes in ((OrderStage.PENDING, 5), (OrderStage.PAID, -1), (OrderStage.CANCELLED, -1)):
            with self.subTest(stage=stage, minutes=minutes):
                self.set_stage(stage)
                self.expire(minutes)
                check_order_stage(self.order.id, self.invoice.id)
                self.assert_stage(self.order, stage)
                self.assert_locked(self.examination, True)
        self.invoice.refresh_from_db()
        self.assertEqual(self.invoice.stage, InvoiceStage.PENDING)

    def test_missing_order_is_ignored(self):
        self.assertIsNone(check_order_stage(self.order.id + 1000, self.invoice.id))
        self.assert_stage(self.order, OrderStage.PENDING)
